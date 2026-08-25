"""
Tableau Extract Reader
======================
Reads rows out of the .hyper extracts bundled inside a .twbx, so a published
model holds real data instead of an empty shell.

This is the Tableau counterpart to `qlik_data_reader.py` and answers the same
contract: `read_workbook_tables(...) -> (data, problems)`, where `data` maps a
table name to {columns, rows, truncated} and `problems` maps a table name to the
reason it could not be read. A table absent from `data` is always present in
`problems`, so a caller can never mistake "not read" for "empty".

Three things make this different from the Qlik side:

1. **The reader is optional.** `tableauhyperapi` is a large native package, and
   a workbook on a live database connection carries no extract to read. It is
   therefore imported lazily, and its absence is reported as a gap rather than
   raised as an error.

2. **Extracts are often denormalised.** Tableau frequently materialises a join
   into one flat table rather than preserving the source tables. When that
   happens the extract's table cannot be matched to the model's tables, and this
   module says so explicitly instead of guessing which columns belong where —
   projecting them apart would duplicate rows wherever the join was one-to-many
   and quietly change every total.

3. **.tde is not readable.** Workbooks saved before Tableau 10.5 carry the older
   .tde format, which has no public reader. It is reported, never approximated.
"""

import os
import zipfile

DEFAULT_MAX_ROWS = 50000


class TableauDataError(RuntimeError):
    """Raised when an extract exists but cannot be read at all."""


def _note(note):
    return note or (lambda _text: None)


def find_extracts(twbx_path):
    """Every extract inside the workbook, as (archive_name, format) pairs."""
    if not zipfile.is_zipfile(twbx_path):
        return []
    found = []
    with zipfile.ZipFile(twbx_path) as archive:
        for name in archive.namelist():
            lowered = name.lower()
            if lowered.endswith(".hyper"):
                found.append((name, "hyper"))
            elif lowered.endswith(".tde"):
                found.append((name, "tde"))
    return found


# Raw data files a .twbx packages when its connection is to a local file and the
# workbook was saved WITHOUT an extract. Tableau bundles the source file itself,
# so these hold the rows even though there is no .hyper anywhere in the archive.
_DELIMITED_SUFFIXES = (".csv", ".tsv", ".txt")
_EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xls")


def find_packaged_data(archive_path):
    """Raw data files bundled in the archive, as (name, kind) pairs.

    Only consulted when there is no extract: an extract is always the better
    source, because it is what Tableau itself would have queried.
    """
    if not zipfile.is_zipfile(archive_path):
        return []
    found = []
    with zipfile.ZipFile(archive_path) as archive:
        for info in archive.infolist():
            if info.is_dir() or not info.file_size:
                continue
            lowered = info.filename.lower()
            if lowered.endswith(_DELIMITED_SUFFIXES):
                found.append((info.filename, "delimited"))
            elif lowered.endswith(_EXCEL_SUFFIXES):
                found.append((info.filename, "excel"))
    return found


def _looks_numeric(text):
    """Whether this cell is a number the file wrote as a number.

    A leading zero is significant -- "01234" is an identifier, not a quantity --
    so it disqualifies the value, as does any letter (which rules out "nan" and
    "inf", both of which float() would otherwise accept).
    """
    body = text.lstrip("-")
    if not body or any(ch.isalpha() for ch in body):
        return False
    if len(body) > 1 and body[0] == "0" and not body.startswith("0."):
        return False
    try:
        float(text)
        return True
    except ValueError:
        return False


def _type_columns(columns, raw_rows):
    """Converts a column to numbers only when every value in it is one.

    Decided per column rather than per cell. Typing cell by cell produced
    columns holding a mix of int and str -- "01234" kept as text next to 90210
    read as an integer -- which is not a schema Parquet or Tabular can express,
    and which would have been resolved by silently discarding one of them.
    """
    typed = [list(row) for row in raw_rows]

    for index in range(len(columns)):
        values = [row[index] for row in typed]
        present = [v for v in values if v is not None and str(v).strip() != ""]
        if not present or not all(_looks_numeric(str(v).strip()) for v in present):
            # Text column: leave every value as written, blanks as null.
            for row in typed:
                text = row[index]
                row[index] = None if text is None or str(text).strip() == "" else str(text)
            continue

        whole = all(float(str(v).strip()).is_integer() and "." not in str(v)
                    for v in present)
        for row in typed:
            text = row[index]
            if text is None or str(text).strip() == "":
                row[index] = None
            else:
                row[index] = int(str(text).strip()) if whole else float(str(text).strip())

    return typed


def _read_delimited_member(archive_path, member, max_rows, note):
    """Reads one packaged delimited file into {columns, rows, truncated}."""
    import csv
    import io as _io

    with zipfile.ZipFile(archive_path) as archive:
        raw = archive.read(member)

    # Tableau writes these out in whatever the source used; UTF-8 with a BOM is
    # the common case and utf-8-sig strips it without corrupting the header.
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TableauDataError("Could not decode %s in any known encoding." % member)

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
        if member.lower().endswith(".tsv"):
            dialect = csv.excel_tab

    reader = csv.reader(_io.StringIO(text), dialect)
    try:
        header = next(reader)
    except StopIteration:
        return None

    columns = [h.strip() or ("Column%d" % (i + 1)) for i, h in enumerate(header)]
    rows, truncated = [], False
    for row in reader:
        if max_rows is not None and len(rows) >= max_rows:
            truncated = True
            break
        # Short and long rows are padded/trimmed to the header, so a ragged file
        # does not shift every later column by one.
        cells = list(row[:len(columns)]) + [None] * (len(columns) - len(row))
        rows.append(cells)

    # Typed after every row is in hand: a column's type is a property of all of
    # its values, not of the first one seen.
    rows = _type_columns(columns, rows)

    note("[extract] %s: %d row(s)%s, %d column(s)"
         % (os.path.basename(member), len(rows),
            " (truncated)" if truncated else "", len(columns)))
    return {"columns": columns, "rows": rows, "truncated": truncated}


def _read_excel_member(archive_path, member, max_rows, note, work_dir):
    """Reads one packaged Excel worksheet, if openpyxl is available."""
    try:
        from openpyxl import load_workbook
    except ImportError as err:
        raise TableauDataError(
            "This workbook packages its data as the Excel file %r rather than as "
            "an extract, and the reader for it is not installed, so no rows were "
            "read.\n\nInstall it with:  pip install openpyxl\n\n"
            "Underlying import error: %s" % (os.path.basename(member), err))

    destination = os.path.join(work_dir, os.path.basename(member))
    with zipfile.ZipFile(archive_path) as archive:
        with archive.open(member) as source, open(destination, "wb") as target:
            target.write(source.read())

    try:
        book = load_workbook(destination, read_only=True, data_only=True)
        sheet = book[book.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return None
        columns = [str(h).strip() if h is not None else ("Column%d" % (i + 1))
                   for i, h in enumerate(header)]
        rows, truncated = [], False
        for row in rows_iter:
            if max_rows is not None and len(rows) >= max_rows:
                truncated = True
                break
            cells = list(row[:len(columns)]) + [None] * (len(columns) - len(row))
            rows.append(list(cells))
        book.close()
    finally:
        try:
            os.remove(destination)
        except OSError:
            pass

    note("[extract] %s: %d row(s)%s, %d column(s)"
         % (os.path.basename(member), len(rows),
            " (truncated)" if truncated else "", len(columns)))
    return {"columns": columns, "rows": rows, "truncated": truncated}


def read_packaged_data(archives, max_rows, work_dir, note):
    """Rows from raw data files packaged in the archives, keyed by file stem."""
    tables, problems = {}, {}
    for archive_path in archives:
        for member, kind in find_packaged_data(archive_path):
            stem = os.path.splitext(os.path.basename(member))[0]
            try:
                if kind == "delimited":
                    payload = _read_delimited_member(archive_path, member, max_rows, note)
                else:
                    payload = _read_excel_member(archive_path, member, max_rows,
                                                 note, work_dir)
            except TableauDataError as err:
                problems[stem] = str(err)
                continue
            except Exception as err:                  # noqa: BLE001
                problems[stem] = "Could not read packaged file %r: %s" % (member, err)
                continue
            if payload and payload["columns"]:
                tables[stem] = payload
    return tables, problems


def _import_hyper_api():
    """Imports tableauhyperapi, or explains precisely why it is unavailable."""
    try:
        from tableauhyperapi import (          # noqa: F401
            Connection, HyperProcess, Telemetry, TableName, CreateMode,
        )
    except ImportError as err:
        raise TableauDataError(
            "The workbook carries a .hyper extract, but the reader for it is not "
            "installed, so no rows were read.\n\n"
            "Install it with:  pip install tableauhyperapi\n\n"
            "It is optional on purpose: workbooks on live database connections "
            "carry no extract, and the package is a large native download that "
            "those migrations never need.\n\n"
            "Underlying import error: %s" % err
        )
    from tableauhyperapi import Connection, HyperProcess, Telemetry, TableName
    return Connection, HyperProcess, Telemetry, TableName


def _extract_to_disk(twbx_path, archive_name, work_dir, note):
    """Unpacks one extract beside the workbook so the Hyper API can open it.

    The API takes a filesystem path, not a stream, so the extract has to be
    materialised. It is written under `work_dir`, which the caller owns and
    cleans up with the rest of the run.
    """
    destination = os.path.join(work_dir, os.path.basename(archive_name))
    with zipfile.ZipFile(twbx_path) as archive:
        with archive.open(archive_name) as source, open(destination, "wb") as target:
            while True:
                chunk = source.read(1024 * 1024)
                if not chunk:
                    break
                target.write(chunk)
    note("[extract] Unpacked %s (%.1f MB)."
         % (os.path.basename(archive_name), os.path.getsize(destination) / 1048576.0))
    return destination


def _read_hyper(hyper_path, max_rows, note):
    """Reads every table in one .hyper file.

    Returns {table_name: {columns, rows, truncated}}.
    """
    Connection, HyperProcess, Telemetry, TableName = _import_hyper_api()

    tables = {}
    # DO_NOT_SEND_USAGE_DATA_TO_TABLEAU: this runs inside someone's migration of
    # their own data, which is not ours to report on.
    with HyperProcess(telemetry=Telemetry.DO_NOT_SEND_USAGE_DATA_TO_TABLEAU) as hyper:
        with Connection(endpoint=hyper.endpoint, database=hyper_path) as connection:
            schemas = connection.catalog.get_schema_names()
            note("[extract] %d schema(s): %s"
                 % (len(schemas), ", ".join(str(s) for s in schemas) or "none"))

            for schema in schemas:
                for table_name in connection.catalog.get_table_names(schema):
                    definition = connection.catalog.get_table_definition(table_name)
                    columns = [str(column.name).strip('"') for column in definition.columns]
                    if not columns:
                        continue

                    # A row cap keeps embedded data inside the request body the
                    # Fabric API accepts. One extra row is requested so a
                    # truncated table can be reported as truncated rather than
                    # silently presented as complete.
                    limit = "" if max_rows is None else " LIMIT %d" % (max_rows + 1)
                    query = "SELECT * FROM %s%s" % (table_name, limit)

                    try:
                        result = connection.execute_list_query(query)
                    except Exception as err:              # noqa: BLE001
                        note("[extract] Could not read %s: %s" % (table_name, err))
                        continue

                    truncated = max_rows is not None and len(result) > max_rows
                    if truncated:
                        result = result[:max_rows]

                    rows = [[_coerce(value) for value in row] for row in result]
                    plain = str(getattr(table_name, "name", table_name)).strip('"')
                    tables[plain] = {
                        "columns": columns,
                        "rows": rows,
                        "truncated": truncated,
                    }
                    note("[extract] %s: %d row(s)%s, %d column(s)"
                         % (plain, len(rows), " (truncated)" if truncated else "",
                            len(columns)))
    return tables


def _coerce(value):
    """Makes a Hyper value safe to serialise into the model's embedded rows."""
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    # date, datetime, Interval, bytes and the rest carry no JSON form. Their
    # string form is what the extract itself states, so nothing is invented.
    return str(value)


def read_workbook_tables(twbx_path, table_names=None, max_rows=DEFAULT_MAX_ROWS,
                         work_dir=None, note=None, extra_archives=None):
    """Reads rows for the named tables out of a workbook's extracts.

    Returns (data, problems), matching the Qlik reader's contract exactly so the
    engine consumes both the same way.
    """
    note = _note(note)
    data, problems = {}, {}
    wanted = [str(n) for n in (table_names or []) if n]

    if not os.path.exists(twbx_path):
        return {}, {"(all tables)": "The workbook file is no longer on disk."}

    # Archives whose extracts belong to this workbook but live outside it --
    # a datasource published separately on the server keeps its rows in its own
    # item, so the workbook itself packages nothing.
    archives = [twbx_path] + [p for p in (extra_archives or []) if os.path.exists(p)]

    extracts = []
    for archive in archives:
        extracts.extend((archive, name, kind) for name, kind in find_extracts(archive))

    if not extracts:
        # No extract does not mean no data. A .twbx whose connection is to a
        # local file and which was saved WITHOUT an extract packages the source
        # file itself -- the rows are sitting in the archive as a CSV or an
        # Excel workbook. Treating that as "live connection, nothing to
        # migrate" threw away data the archive was carrying all along.
        work_dir = work_dir or os.path.dirname(os.path.abspath(twbx_path))
        packaged, packaged_problems = read_packaged_data(
            archives, max_rows, work_dir, note)

        if packaged:
            note("[extract] No extract, but the workbook packages its source "
                 "data: %s." % ", ".join(sorted(packaged)))
            return _match_tables(packaged, wanted, note, packaged_problems)

        reason = ("This workbook carries no extract and packages no data file, so "
                  "it holds no rows — its data lives in the connected database. "
                  "The model was built with its structure only.")
        note("[extract] %s" % reason)
        problems = {name: reason for name in (wanted or ["(all tables)"])}
        problems.update(packaged_problems)
        return {}, problems

    hyper_files = [(archive, name) for archive, name, kind in extracts if kind == "hyper"]
    tde_files = [(archive, name) for archive, name, kind in extracts if kind == "tde"]

    if tde_files and not hyper_files:
        reason = ("This workbook's extract is in the legacy .tde format (Tableau "
                  "10.4 and earlier), which has no public reader. Re-save the "
                  "workbook in a current Tableau version to convert it to .hyper, "
                  "then migrate again. No rows were read.")
        note("[extract] %s" % reason)
        return {}, {name: reason for name in (wanted or ["(all tables)"])}

    work_dir = work_dir or os.path.dirname(os.path.abspath(twbx_path))
    found = {}

    for source_archive, archive_name in hyper_files:
        try:
            hyper_path = _extract_to_disk(source_archive, archive_name, work_dir, note)
        except (OSError, zipfile.BadZipFile) as err:
            note("[extract] Could not unpack %s: %s" % (archive_name, err))
            continue

        try:
            found.update(_read_hyper(hyper_path, max_rows, note))
        except TableauDataError as err:
            # The reader is missing entirely; no later extract will fare better.
            note("[extract] %s" % err)
            return {}, {name: str(err) for name in (wanted or ["(all tables)"])}
        except Exception as err:                          # noqa: BLE001
            note("[extract] Failed reading %s: %s" % (archive_name, err))
        finally:
            try:
                os.remove(hyper_path)
            except OSError:
                pass

    if not found:
        reason = "The workbook's extract could not be read, so no rows are available."
        return {}, {name: reason for name in (wanted or ["(all tables)"])}

    return _match_tables(found, wanted, note, problems)


def _match_tables(found, wanted, note, problems=None):
    """Pairs the tables a source yielded with the tables the model declares.

    Shared by the extract path and the packaged-data-file path: both end up
    holding {name: payload} and needing the same reconciliation against the
    model's own table names.
    """
    data = {}
    problems = dict(problems or {})
    # Match the source's tables to the model's, case-insensitively.
    by_lower = {name.lower(): name for name in found}

    # A single-table workbook whose extract holds a single table is that table,
    # whatever the two happen to be called. Tableau names an extract's table
    # after the extract ("Extract"), not after the source relation, so a
    # one-to-one workbook otherwise matched nothing and migrated with no rows at
    # all. This is identity, not guesswork: with one table on each side there is
    # no other pairing to make. Anything less clear-cut is still reported.
    if len(found) == 1 and len(wanted) == 1 and wanted[0].lower() not in by_lower:
        only_extract = next(iter(found))
        note("[extract] The extract's only table is %r and the model's only table "
             "is %r; reading one as the other." % (only_extract, wanted[0]))
        return {wanted[0]: found[only_extract]}, {}

    for name in (wanted or list(found)):
        actual = by_lower.get(name.lower())
        if actual:
            data[name] = found[actual]
            continue

        # The usual reason for a miss: Tableau materialised the join into one
        # flat table. Splitting it back apart would duplicate rows wherever the
        # join was one-to-many, so it is reported rather than attempted.
        problems[name] = (
            "The extract carries no table named %r. It holds: %s. Tableau often "
            "materialises a join into a single flat table, in which case the "
            "source tables are not present in it individually. Those rows were "
            "not split apart by guesswork, which would multiply rows wherever the "
            "join was one-to-many and change every total; the extract's own "
            "table is migrated whole instead."
            % (name, ", ".join(sorted(found)) or "nothing")
        )

    # Extract tables the model never named are carried over under their own
    # names rather than discarded. The engine adds them to the model as tables
    # in their own right and reports them, which is what the Qlik path already
    # does for tables its script parser did not predict.
    #
    # This is what makes a denormalised extract migrate at all: the model's
    # normalised tables stay empty and say why, while the rows themselves still
    # reach Fabric intact and unduplicated.
    claimed = {w.lower() for w in (wanted or [])}
    unused = [name for name in found if name.lower() not in claimed
              and name not in data]
    for name in unused:
        data[name] = found[name]
    if wanted and unused:
        note("[extract] The extract holds %d table(s) the model does not name; "
             "migrating them under their own names: %s"
             % (len(unused), ", ".join(sorted(unused))))

    return data, problems
